"""Audit user-provided education workbooks without publishing observations or changing originals."""

import argparse
import hashlib
import json
import re
from collections import Counter
from decimal import Decimal, localcontext
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
COMMON_LEVELS = {"TK", "SD", "SMP", "SMA", "SMK", "SLB"}
TOLERANCE = Decimal("0.00000001")


def number(value, unit):
    if value is None or value in ("", "-", "…"):
        return None, "missing"
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None, "invalid_type"
    result = Decimal(str(value))
    if (
        not result.is_finite()
        or result < 0
        or (unit == "percent" and result > 100)
        or (unit == "school_count" and result != result.to_integral_value())
    ):
        return None, "invalid_range"
    return result, None


def normalize_value(value, source, number_format):
    scale = source["storage_scale"]
    if scale == "fraction_with_percent_format":
        if source["unit"] != "percent" or number_format != "0.00%":
            raise ValueError("Percentage storage format changed")
        result, issue = number(value, "percent")
        if result is not None and result > 1:
            return None, "invalid_range"
        return (result * 100 if result is not None else None), issue
    if (scale, source["unit"]) not in (("percentage_points", "percent"), ("count", "school_count")):
        raise ValueError("Unsupported declared storage scale")
    return number(value, source["unit"])


def inspect_workbook(raw_dir, source):
    path = raw_dir / source["filename"]
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if digest != source["sha256"]:
        raise ValueError(f"{path.name}: source hash changed; review a new version")
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = book["Data"]
        rows = list(sheet.values)
        cell_rows = list(sheet.iter_rows())
        metadata = {row[0]: row[1] for row in book["Metadata"].values if row[0]}
        if metadata.get("Nama") != source["title"] or rows[0][0] != source["title"]:
            raise ValueError(f"{path.name}: metadata/title mismatch")
        unit = {"Persen": "percent", "Unit": "school_count"}.get(
            str(metadata.get("Satuan")).strip()
        )
        if unit != source["unit"] or "2025/2026" not in source["title"]:
            raise ValueError(f"{path.name}: unit or period mismatch")
        header = next(
            i for i, row in enumerate(rows[:10]) if row[:3] == ("No", "Kode Wilayah", "Wilayah")
        )
        status = next(
            i
            for i in range(header + 1, min(header + 4, len(rows)))
            if rows[i][3:6] == ("Negeri", "Swasta", "Total")
        )
        level_row = rows[status - 1] if "(TK)" in str(rows[status - 1][3]) else rows[header]
        groups = {}
        for column, label in enumerate(level_row[3:], 3):
            if label is None:
                continue
            match = re.search(r"\(([^)]+)\)", str(label))
            level = match.group(1) if match else str(label)
            if level in groups or rows[status][column : column + 3] != (
                "Negeri",
                "Swasta",
                "Total",
            ):
                raise ValueError(f"{path.name}: unsupported level/ownership layout")
            groups[level] = column
        if not COMMON_LEVELS <= set(groups):
            raise ValueError(f"{path.name}: common education levels are missing")
        values, names, issues, excluded = {}, {}, [], []
        counts = Counter()
        total_mismatches = []
        used = {column + offset for column in groups.values() for offset in range(3)}
        for index, row in enumerate(rows[status + 1 :], status + 2):
            if all(value is None for value in row):
                continue
            if row[2] in ("Indonesia", "Luar Negeri"):
                excluded.append(row[2])
                continue
            code = str(row[1])
            if (
                not isinstance(row[2], str)
                or not row[2].startswith("Provinsi ")
                or not re.fullmatch(r"[0-9]{2}", code)
            ):
                raise ValueError(f"{path.name}: unrecognized region at row {index}")
            if code in names:
                raise ValueError(f"{path.name}: duplicate region code {code}")
            names[code] = row[2]
            if any(
                value is not None for column, value in enumerate(row[3:], 3) if column not in used
            ):
                raise ValueError(f"{path.name}: data outside recognized header groups")
            for level, start in groups.items():
                for offset, ownership in enumerate(("Negeri", "Swasta", "Total")):
                    raw_value = row[start + offset]
                    value, issue = normalize_value(
                        raw_value, source, cell_rows[index - 1][start + offset].number_format
                    )
                    coordinate = f"{get_column_letter(start + offset + 1)}{index}"
                    if isinstance(row[start + offset], str) and row[start + offset].startswith("="):
                        issue = "formula_not_allowed"
                    counts["observations"] += 1
                    if issue:
                        counts[issue] += 1
                        issues.append({"cell": coordinate, "issue": issue})
                    else:
                        counts["numeric"] += 1
                        counts["zeros"] += value == 0
                    values[(code, level, ownership)] = value
                if unit == "school_count":
                    parts = [
                        values[(code, level, ownership)]
                        for ownership in ("Negeri", "Swasta", "Total")
                    ]
                    if (
                        all(value is not None for value in parts)
                        and parts[0] + parts[1] != parts[2]
                    ):
                        total_mismatches.append({"source_region_code": code, "level": level})
        report = {
            "id": source["id"],
            "filename": path.name,
            "sha256": digest,
            "bytes": path.stat().st_size,
            "title": source["title"],
            "unit": unit,
            "sheets": book.sheetnames,
            "storage_scale": source["storage_scale"],
            "audit_value_unit": "percentage_points" if unit == "percent" else "school_count",
            "columns_in_file": sheet.max_column,
            "education_levels": list(groups),
            "province_rows": len(names),
            "duplicate_region_codes": [],
            "excluded_aggregates": excluded,
            "quality": dict(counts),
            "issues": issues,
            "ownership_total_mismatches": total_mismatches,
            "metadata_license_field_present": any(
                re.search(r"lisensi|license|ketentuan", str(key), re.I) for key in metadata
            ),
            "matches_prior_source": source["matches_prior_source"],
        }
        return report, values, names
    finally:
        book.close()


def reconcile(counts, percentages, denominator):
    outcomes = Counter()
    sd = Counter()
    mismatches = []
    max_difference = Decimal(0)
    with localcontext() as context:
        context.prec = 50
        for key, numerator in counts.items():
            if key[1] not in COMMON_LEVELS:
                continue
            total, percentage = denominator.get(key), percentages.get(key)
            if numerator is None or total is None or percentage is None:
                outcome = "missing_or_invalid"
            elif total == 0:
                outcome = "zero_denominator"
            else:
                difference = abs(100 * numerator / total - percentage)
                max_difference = max(max_difference, difference)
                outcome = "match" if difference <= TOLERANCE else "mismatch"
                if outcome == "mismatch":
                    mismatches.append(
                        {"source_region_code": key[0], "level": key[1], "ownership": key[2]}
                    )
            outcomes[outcome] += 1
            if key[1:] == ("SD", "Total"):
                sd[outcome] += 1
    return {
        "tolerance_percentage_points": str(TOLERANCE),
        "outcomes": dict(outcomes),
        "sd_total_outcomes": dict(sd),
        "max_absolute_difference_percentage_points": str(max_difference),
        "mismatch_keys": mismatches,
    }


def audit(raw_dir, manifest):
    reports, values, names = [], {}, {}
    for source in manifest["sources"]:
        report, values[source["id"]], names[source["id"]] = inspect_workbook(raw_dir, source)
        reports.append(report)
    union = set.union(*(set(rows) for rows in names.values()))
    intersection = set.intersection(*(set(rows) for rows in names.values()))
    denominator = values["institutions_count"]
    bounds = {}
    for source in (
        "internet_count",
        "computers_count",
        "water_count",
        "electricity_count",
        "handwashing_count",
    ):
        eligible = [(key, value, denominator.get(key)) for key, value in values[source].items()]
        bounds[source] = {
            "checked": sum(value is not None and total is not None for _, value, total in eligible),
            "exceeds_denominator": [
                {"source_region_code": key[0], "level": key[1], "ownership": key[2]}
                for key, value, total in eligible
                if value is not None and total is not None and value > total
            ],
        }
    name_mismatches = {
        source: sorted(
            code for code in intersection if rows[code] != names["institutions_count"][code]
        )
        for source, rows in names.items()
    }
    return {
        "checked_at": manifest["checked_at"],
        "release_status": "not_approved",
        "downloaded_at": None,
        "raw_directory": manifest["local_directory"],
        "sources": reports,
        "province_join": {
            "union": len(union),
            "intersection": len(intersection),
            "name_mismatches": name_mismatches,
        },
        "count_bounds": bounds,
        "percentage_reconciliation": {
            indicator: reconcile(
                values[f"{indicator}_count"], values[f"{indicator}_percent"], denominator
            )
            for indicator in ("internet", "computers", "water")
        },
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-dir", type=Path, default=ROOT / "data/datamentah")
    parser.add_argument(
        "--output", type=Path, default=ROOT / "data/metadata/additional-source-audit.json"
    )
    args = parser.parse_args()
    manifest = json.loads(
        (ROOT / "data/metadata/additional-source-candidates.json").read_text("utf-8")
    )
    result = audit(args.raw_dir, manifest)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", "utf-8")
    print(
        json.dumps(
            {
                "files": len(result["sources"]),
                "province_intersection": result["province_join"]["intersection"],
                "percentage_reconciliation": {
                    indicator: {
                        key: value for key, value in report.items() if key != "mismatch_keys"
                    }
                    for indicator, report in result["percentage_reconciliation"].items()
                },
            }
        )
    )


if __name__ == "__main__":
    main()
