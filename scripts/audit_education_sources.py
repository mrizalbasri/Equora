"""Inspect pinned source workbooks locally; never publish observations or approve a release."""

import argparse
import hashlib
import json
import math
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
LEVELS = ("TK", "SD", "SMP", "SMA", "SMK", "SLB")


def inspect_source(path, source):
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if digest != source["sha256"]:
        raise ValueError(
            f"{path.name}: SHA256 changed; audit a new source version explicitly"
        )
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = book["Data"]
        rows = [
            (i, row)
            for i, row in enumerate(sheet.values, 1)
            if any(v is not None for v in row)
        ]
        header = next(
            i for i, row in rows if row[:3] == ("No", "Kode Wilayah", "Wilayah")
        )
        provinces = [
            (i, r)
            for i, r in rows
            if isinstance(r[2], str) and r[2].startswith("Provinsi ")
        ]
        if not provinces or sheet.max_column != 21:
            raise ValueError(f"{path.name}: unexpected province layout")
        first = provinces[0][0]
        headers = {i: row for i, row in rows if header <= i < first}
        level_row = next(
            row for row in headers.values() if row[3] == "Taman Kanak-kanak (TK)"
        )
        status_row = headers[first - 1]
        for index, level in enumerate(LEVELS):
            start = 3 + index * 3
            if f"({level})" not in str(level_row[start]):
                raise ValueError(f"{path.name}: changed education order at {level}")
            if status_row[start : start + 3] != ("Negeri", "Swasta", "Total"):
                raise ValueError(f"{path.name}: changed ownership order at {level}")
        codes = [str(row[1]) for _, row in provinces]
        if not all(code.isdigit() and len(code) == 2 for code in codes):
            raise ValueError(
                f"{path.name}: province codes require explicit mapping review"
            )
        stats = Counter()
        missing_tokens = Counter()
        complete_by_level = {level: set() for level in LEVELS}
        sd_codes = set()
        for _, row in provinces:
            for index, value in enumerate(row[3:21]):
                stats["observations"] += 1
                if value is None or value in ("", "-", "…"):
                    stats["missing"] += 1
                    missing_tokens[str(value)] += 1
                elif isinstance(value, bool) or not isinstance(value, (int, float)):
                    stats["invalid_type"] += 1
                elif (
                    not math.isfinite(value)
                    or value < 0
                    or (source["unit"] == "percent" and value > 100)
                    or (source["unit"] == "school_count" and value != int(value))
                ):
                    stats["invalid_range"] += 1
                else:
                    stats["numeric"] += 1
                    stats["zeros"] += value == 0
                    if index % 3 == 2:
                        complete_by_level[LEVELS[index // 3]].add(str(row[1]))
                        if index == 5:
                            sd_codes.add(str(row[1]))
        excluded = [
            row[2]
            for i, row in rows
            if i >= first and not str(row[2]).startswith("Provinsi ")
        ]
        metadata = [
            row for row in book["Metadata"].values if any(v is not None for v in row)
        ]
        report = {
            "id": source["id"],
            "filename": path.name,
            "sha256": digest,
            "bytes": path.stat().st_size,
            "retrieved_at": datetime.fromtimestamp(
                path.stat().st_mtime, UTC
            ).isoformat(),
            "title": sheet["A1"].value,
            "sheets": book.sheetnames,
            "header_rows": list(headers),
            "first_data_row": first,
            "province_rows": len(provinces),
            "unique_province_codes": len(set(codes)),
            "duplicate_province_codes": [
                code for code, n in Counter(codes).items() if n > 1
            ],
            "excluded_aggregate_rows": excluded,
            "province_observation_quality": {
                key: stats[key]
                for key in (
                    "observations",
                    "numeric",
                    "zeros",
                    "missing",
                    "invalid_type",
                    "invalid_range",
                )
            },
            "missing_tokens": dict(missing_tokens),
            "nonmissing_total_by_education_level": {
                k: len(v) for k, v in complete_by_level.items()
            },
            "license_field_present": any(
                any(
                    term in str(row[0]).lower()
                    for term in ("lisensi", "license", "ketentuan")
                )
                for row in metadata
            ),
        }
        preview = {"header_rows": headers, "first_ten_province_rows": provinces[:10]}
        return report, set(codes), sd_codes, preview
    finally:
        book.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--raw-dir", type=Path, default=ROOT / "data/raw/source-audit-2026-10-19"
    )
    parser.add_argument(
        "--output", type=Path, default=ROOT / "data/metadata/source-audit.json"
    )
    args = parser.parse_args()
    manifest = json.loads(
        (ROOT / "data/metadata/source-candidates.json").read_text("utf-8")
    )
    reports, code_sets, sd_sets = [], [], []
    for source in manifest["sources"]:
        path = args.raw_dir / source["filename"]
        report, codes, sd_codes, preview = inspect_source(path, source)
        reports.append(report)
        code_sets.append(codes)
        sd_sets.append(sd_codes)
        path.with_suffix(".preview.json").write_text(
            json.dumps(preview, indent=2), "utf-8"
        )
    union, intersection = set.union(*code_sets), set.intersection(*code_sets)
    output = {
        "release_status": "not_approved",
        "sources": reports,
        "province_join": {
            "union": len(union),
            "intersection": len(intersection),
            "matched_fraction": len(intersection) / len(union),
            "unmatched_by_source": {
                r["id"]: sorted(union - codes)
                for r, codes in zip(reports, code_sets, strict=True)
            },
            "sd_total_nonmissing_intersection": len(set.intersection(*sd_sets)),
        },
        "limitations": [
            "Source province codes have not been mapped to a reviewed boundary or regional registry.",
            "No denominator is provided for internet/computer counts.",
            "No license/terms approval, cost assumptions, or release approval is implied.",
        ],
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(output, ensure_ascii=False, indent=2) + "\n", "utf-8"
    )
    print(
        json.dumps(
            {"output": str(args.output), "province_join": output["province_join"]}
        )
    )


if __name__ == "__main__":
    main()
