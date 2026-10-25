"""Explicit local-review snapshot from the user's ten hash-pinned workbooks."""

import hashlib
import json
import re
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.etl.additional_education import (
    PROFILES,
    PUBLISHER,
    AdditionalSourceMetadata,
    audited_additional_source,
    stage_additional_workbook,
)
from app.etl.education import LEVELS, Issue, Observation, SourceMetadata, StagingResult, _number

ADAPTER = "education-header-and-scale-v2"
DEFINITIONS = {
    "institutions_count": (
        "Jumlah satuan pendidikan pada wilayah, jenjang dan tahun ajaran yang sama."
    ),
    "internet_count": "Jumlah sekolah dengan akses internet untuk pengajaran.",
    "computers_count": "Jumlah sekolah yang memiliki sedikitnya lima komputer untuk pengajaran.",
    "electricity_count": "Jumlah sekolah dengan akses listrik.",
    "water_count": "Jumlah sekolah dengan air minum layak, disediakan sekolah atau dibawa siswa.",
    "handwashing_count": "Jumlah sekolah dengan sarana cuci tangan, sabun dan air mengalir.",
    "internet_percent": "Persentase sekolah dengan akses internet untuk pengajaran.",
    "computers_percent": "Persentase sekolah dengan sedikitnya lima komputer untuk pengajaran.",
    "water_percent": "Persentase sekolah dengan akses air minum layak.",
    "sanitation_percent": "Persentase sekolah dengan toilet peserta didik laki-laki dan perempuan.",
}
CORE = ("sanitation_percent", "water_percent", "internet_percent", "computers_percent")


def read_source(path: Path, record: dict, publisher: str = "Pusdatin Kemendikdasmen"):
    if record["id"] in PROFILES:
        metadata = AdditionalSourceMetadata(
            source_id=record["id"],
            title=record["title"],
            publisher=PUBLISHER,
            source_url=record["source_url"],
            sha256=record["sha256"],
            year=2025,
            academic_year="2025/2026",
            unit=record["unit"],
            retrieved_at=record["downloaded_at"],
            license_status="not_verified_for_public_release",
            storage_scale=record["storage_scale"],
            layout=PROFILES[record["id"]][2],
        )
        results = {
            level: stage_additional_workbook(path, metadata, education_level=level)
            for level in LEVELS
        }
        if any(
            result.status != "staged" or result.province_rows != 38 for result in results.values()
        ):
            raise ValueError(
                "Additional source has invalid Total slices or incomplete province coverage; "
                "stage it for a detailed report"
            )
        return results
    content = path.read_bytes()
    if len(content) > 5 * 1024 * 1024 or hashlib.sha256(content).hexdigest() != record["sha256"]:
        raise ValueError("Source size or SHA256 differs from the audited file")
    expected_unit = "school_count" if record["storage_scale"] == "count" else "percent"
    if (
        record["storage_scale"]
        not in ("count", "percentage_points", "fraction_with_percent_format")
        or record["unit"] != expected_unit
    ):
        raise ValueError("Unknown or incompatible storage scale")
    book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        sheet = book["Data"]
        if sheet.max_row > 1000 or sheet.max_column > 60:
            raise ValueError("Unexpected source dimensions")
        rows = list(sheet.iter_rows())
        if rows[0][0].value != record["title"]:
            raise ValueError("Source title changed")
        header = next(
            i
            for i, row in enumerate(rows[:10])
            if tuple(c.value for c in row[:3]) == ("No", "Kode Wilayah", "Wilayah")
        )
        status = next(
            i
            for i in range(header + 1, header + 4)
            if tuple(c.value for c in rows[i][3:6]) == ("Negeri", "Swasta", "Total")
        )
        level_row = rows[status - 1] if "(TK)" in str(rows[status - 1][3].value) else rows[header]
        columns = {}
        for column, cell in enumerate(level_row):
            for level in LEVELS:
                if f"({level})" in str(cell.value):
                    if level in columns or tuple(
                        c.value for c in rows[status][column : column + 3]
                    ) != ("Negeri", "Swasta", "Total"):
                        raise ValueError("Ambiguous education/ownership columns")
                    columns[level] = column + 2
        if set(columns) != set(LEVELS):
            raise ValueError("Education columns are incomplete")
        metadata = SourceMetadata(
            source_id=record["id"],
            title=record["title"],
            publisher=publisher,
            source_url=record["source_url"],
            sha256=record["sha256"],
            year=2025,
            academic_year="2025/2026",
            unit=record["unit"],
            retrieved_at=record.get("downloaded_at"),
            license_status="not_verified_for_public_release",
        )
        observations = {level: [] for level in LEVELS}
        issues = {level: [] for level in LEVELS}
        names = {}
        for index, row in enumerate(rows[status + 1 :], status + 2):
            if row[2].value in ("Indonesia", "Luar Negeri") or all(c.value is None for c in row):
                continue
            code, name = str(row[1].value), row[2].value
            if (
                not re.fullmatch(r"[0-9]{2}", code)
                or code in names
                or not str(name).startswith("Provinsi ")
            ):
                raise ValueError("Unknown/duplicate province row")
            names[code] = name
            for level, column in columns.items():
                cell = row[column]
                coordinate = f"{get_column_letter(column + 1)}{index}"
                value, issue = _number(cell.value, metadata.unit)
                if cell.data_type == "f":
                    raise ValueError("Formula in a selected Total cell requires review")
                if record["storage_scale"] == "fraction_with_percent_format":
                    if cell.number_format != "0.00%" or (value is not None and value > 1):
                        raise ValueError("Percentage fraction/format changed")
                    if value is not None:
                        value *= 100
                if issue and issue != "missing":
                    raise ValueError(f"Invalid observation at {coordinate}")
                if issue:
                    issues[level].append(Issue(code=issue, severity="warning", cell=coordinate))
                observations[level].append(
                    Observation(
                        source_region_code=code,
                        source_region_name=name,
                        value=value,
                        source_cell=coordinate,
                    )
                )
        if len(names) != 38:
            raise ValueError("Local review expects 38 unique provinces")
        results = {}
        for level in LEVELS:
            missing = sum(o.value is None for o in observations[level])
            results[level] = StagingResult(
                adapter_version=ADAPTER,
                status="staged",
                source=metadata,
                dataset_version_id=f"{metadata.source_id}:sha256:{metadata.sha256}",
                education_level=level,
                observations=tuple(observations[level]),
                issues=tuple(issues[level]),
                excluded_aggregates=("Luar Negeri", "Indonesia"),
                province_rows=38,
                numeric_values=38 - missing,
                missing_values=missing,
                error_count=0,
            )
        return results
    finally:
        book.close()


def build_release(data_dir: Path):
    manifest = json.loads(
        (data_dir / "metadata/additional-source-candidates.json").read_text("utf-8")
    )
    audit = json.loads((data_dir / "metadata/additional-source-audit.json").read_text("utf-8"))
    sources, scopes, staged = [], {level: {} for level in LEVELS}, []
    for record in manifest["sources"]:
        if record["id"] in PROFILES:
            audited_additional_source(record["id"], manifest, audit)
        results = read_source(data_dir / "datamentah" / record["filename"], record)
        sample = results["SD"]
        sources.append(
            {
                **sample.source.model_dump(mode="json"),
                "definition": DEFINITIONS[record["id"]],
                "storage_scale": record["storage_scale"],
                "quality": {
                    level: {"missing": result.missing_values, "numeric": result.numeric_values}
                    for level, result in results.items()
                },
            }
        )
        for level, result in results.items():
            staged.append(result)
            for observation in result.observations:
                region = scopes[level].setdefault(
                    observation.source_region_code,
                    {
                        "code": observation.source_region_code,
                        "name": observation.source_region_name,
                        "values": {},
                        "cells": {},
                    },
                )
                if region["name"] != observation.source_region_name:
                    raise ValueError("Province names differ across sources")
                region["values"][record["id"]] = (
                    str(observation.value) if observation.value is not None else None
                )
                region["cells"][record["id"]] = observation.source_cell
    payload = {
        "status": "local_review",
        "year": 2025,
        "academic_year": "2025/2026",
        "ownership": "Total",
        "adapter_version": ADAPTER,
        "region_namespace": "publisher_source_code",
        "sources": sources,
        "scopes": {
            level: list(sorted(regions.values(), key=lambda row: row["code"]))
            for level, regions in scopes.items()
        },
        "limitations": [
            "Data lokal untuk review; belum release publik.",
            "Tanggal unduh pengguna dan lisensi redistribusi belum diverifikasi.",
            "Nilai agregat provinsi, bukan kondisi setiap sekolah.",
        ],
    }
    fingerprint = hashlib.sha256(json.dumps(payload, sort_keys=True).encode()).hexdigest()
    payload["id"] = f"local-2025-{fingerprint[:16]}"
    return payload, staged
