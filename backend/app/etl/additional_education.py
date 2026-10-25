"""Explicit contracts for the six additions audited in PR #25; Total slices only."""

import re
from pathlib import Path
from typing import Literal

from openpyxl.utils import get_column_letter
from pydantic import model_validator

from app.etl.education import (
    LEVELS,
    MAX_ROWS,
    EducationLevel,
    Issue,
    Observation,
    SourceMetadata,
    StagingResult,
    _number,
    open_verified_workbook,
)

ADAPTER_VERSION = "education-additional-total-staging-v1"
PUBLISHER = "Pusat Data dan Teknologi Informasi, Kemendikdasmen"
DENOMINATOR_LEVELS = ("TK", "KB", "TPA", "SPS", "Kursus", "PKBM", "SKB", *LEVELS[1:])
PROFILES = {
    "institutions_count": ("school_count", "count", "institutions_12_groups"),
    "electricity_count": ("school_count", "count", "indicator_6_groups"),
    "water_count": ("school_count", "count", "indicator_6_groups"),
    "handwashing_count": ("school_count", "count", "indicator_6_groups"),
    "internet_percent": ("percent", "fraction_with_percent_format", "indicator_6_groups"),
    "computers_percent": ("percent", "fraction_with_percent_format", "indicator_6_groups"),
}


class AdditionalSourceMetadata(SourceMetadata):
    storage_scale: Literal["count", "fraction_with_percent_format"]
    layout: Literal["institutions_12_groups", "indicator_6_groups"]

    @model_validator(mode="after")
    def explicit_profile(self):
        if PROFILES.get(self.source_id) != (self.unit, self.storage_scale, self.layout):
            raise ValueError("Unsupported additional source or changed unit/storage scale/layout")
        return self


class AdditionalStagingResult(StagingResult):
    adapter_version: str = ADAPTER_VERSION
    source: AdditionalSourceMetadata
    available_levels: tuple[str, ...]
    unsupported_levels: tuple[str, ...]
    unimported_ownership: tuple[str, ...] = ("Negeri", "Swasta")
    # These cells were checked but do not belong to the requested Total slice.
    unimported_issues: tuple[Issue, ...]


def audited_additional_source(source_id: str, manifest: dict, audit: dict):
    if source_id not in PROFILES:
        raise ValueError(
            "Unsupported additional source; use the original adapter for prior sources"
        )
    records = [row for row in manifest["sources"] if row["id"] == source_id]
    reports = [row for row in audit["sources"] if row["id"] == source_id]
    if len(records) != 1 or len(reports) != 1:
        raise ValueError("Additional source must occur once in both manifest and audit")
    record, report = records[0], reports[0]
    for key in ("sha256", "filename", "title", "unit", "storage_scale"):
        if record[key] != report[key]:
            raise ValueError(f"Additional manifest and audit disagree on {key}")
    expected_levels = DENOMINATOR_LEVELS if source_id == "institutions_count" else LEVELS
    if tuple(report["education_levels"]) != expected_levels:
        raise ValueError("Audited education layout changed")
    if report["columns_in_file"] != (53 if source_id == "institutions_count" else 21):
        raise ValueError("Audited workbook dimensions changed")
    if report["province_rows"] != 38:
        raise ValueError("Audited province coverage changed")
    return AdditionalSourceMetadata(
        source_id=source_id,
        title=record["title"],
        publisher=PUBLISHER,
        source_url=record["source_url"],
        sha256=record["sha256"],
        year=int(manifest["academic_year"].split("/")[0]),
        academic_year=manifest["academic_year"],
        unit=record["unit"],
        retrieved_at=record["downloaded_at"],
        license_status="not_verified_for_public_release",
        storage_scale=record["storage_scale"],
        layout=PROFILES[source_id][2],
    )


def _cell_value(cell, source):
    if cell.data_type == "f":
        return None, "formula_not_allowed"
    value, issue = _number(cell.value, source.unit)
    if issue == "missing":
        return value, issue
    if source.storage_scale == "fraction_with_percent_format":
        if cell.number_format != "0.00%":
            return None, "storage_format_changed"
        if value is not None and value > 1:
            return None, "invalid_range"
        return (value * 100 if value is not None else None), issue
    if "%" in (cell.number_format or ""):
        return None, "storage_format_changed"
    return value, issue


def _columns(rows, expected):
    header = next(
        (
            i
            for i, row in enumerate(rows[:10])
            if tuple(c.value for c in row[:3]) == ("No", "Kode Wilayah", "Wilayah")
        ),
        None,
    )
    if header is None:
        raise ValueError("Unrecognized source header")
    status = next(
        (
            i
            for i in range(header + 1, min(header + 4, len(rows)))
            if tuple(c.value for c in rows[i][3:6]) == ("Negeri", "Swasta", "Total")
        ),
        None,
    )
    if status is None:
        raise ValueError("Unrecognized ownership header")
    candidates = []
    for row in rows[header:status]:
        groups = {}
        for column, cell in enumerate(row[3:], 3):
            if cell.value is None:
                continue
            label = str(cell.value)
            match = re.search(r"\(([^)]+)\)", label)
            level = match.group(1) if match else label
            if level not in expected or level in groups or column % 3 != 0:
                groups = {}
                break
            groups[level] = column
        if set(groups) == set(expected):
            candidates.append(groups)
    if len(candidates) != 1:
        raise ValueError("Unsupported, incomplete or ambiguous education header groups")
    groups = candidates[0]
    used = {start + offset for start in groups.values() for offset in range(3)}
    if used != set(range(3, 3 + 3 * len(expected))):
        raise ValueError("Education groups must be contiguous")
    for start in groups.values():
        if tuple(c.value for c in rows[status][start : start + 3]) != ("Negeri", "Swasta", "Total"):
            raise ValueError("Ownership order changed")
    if any(c.value is not None for i, c in enumerate(rows[status]) if i >= 3 and i not in used):
        raise ValueError("Unsupported ownership header outside recognized groups")
    return status, groups, used


def stage_additional_workbook(
    path: Path, source: AdditionalSourceMetadata, *, education_level: EducationLevel
) -> AdditionalStagingResult:
    if education_level not in LEVELS:
        raise ValueError("Select one supported education level explicitly; PAUD/Dikmas unsupported")
    # Revalidate even if the caller used model_copy/model_construct.
    source = AdditionalSourceMetadata.model_validate(source.model_dump())
    expected = DENOMINATOR_LEVELS if source.layout == "institutions_12_groups" else LEVELS
    book = open_verified_workbook(path, source)
    try:
        if not {"Data", "Metadata"} <= set(book.sheetnames):
            raise ValueError("Workbook requires Data and Metadata sheets")
        sheet = book["Data"]
        if (
            sheet.max_column != (53 if source.layout == "institutions_12_groups" else 21)
            or sheet.max_row > MAX_ROWS
        ):
            raise ValueError("Unexpected additional workbook dimensions")
        rows = list(sheet.iter_rows())
        metadata = {row[0]: row[1] for row in book["Metadata"].values if row[0] and len(row) > 1}
        if not rows or rows[0][0].value != source.title or metadata.get("Nama") != source.title:
            raise ValueError("Workbook title/metadata mismatch")
        if (
            str(metadata.get("Satuan", "")).strip()
            != ("Persen" if source.unit == "percent" else "Unit")
            or source.academic_year not in source.title
        ):
            raise ValueError("Workbook unit or academic year mismatch")
        status, groups, used = _columns(rows, expected)
        observations, issues, unimported, aggregates = [], [], [], []
        if source.retrieved_at is None:
            issues.append(Issue(code="retrieved_at_unknown", severity="warning"))
        issues.append(Issue(code="license_not_verified_for_public_release", severity="warning"))
        if source.source_id in ("electricity_count", "handwashing_count"):
            issues.append(Issue(code="context_only_not_mcda_indicator", severity="warning"))
        codes = set()
        province_rows = 0
        for row_number, row in enumerate(rows[status + 1 :], status + 2):
            if all(c.value is None for c in row):
                continue
            if any(c.value is not None for i, c in enumerate(row) if i >= 3 and i not in used):
                raise ValueError(f"Data outside recognized header groups at row {row_number}")
            name, code = row[2].value, row[1].value
            if name in ("Indonesia", "Luar Negeri"):
                aggregates.append(name)
                continue
            if (
                not isinstance(name, str)
                or not name.startswith("Provinsi ")
                or type(code) not in (str, int)
                or not re.fullmatch(r"[0-9]{2}", str(code))
            ):
                issues.append(Issue(code="invalid_region", severity="error", cell=f"B{row_number}"))
                continue
            province_rows += 1
            code = str(code)
            if code in codes:
                issues.append(
                    Issue(code="duplicate_region", severity="error", cell=f"B{row_number}")
                )
                continue
            codes.add(code)
            for level, start in groups.items():
                for offset in range(3):
                    column = start + offset
                    value, issue = _cell_value(row[column], source)
                    coordinate = f"{get_column_letter(column + 1)}{row_number}"
                    selected = level == education_level and offset == 2
                    if issue:
                        finding = Issue(
                            code=issue,
                            severity="warning" if issue == "missing" else "error",
                            cell=coordinate,
                        )
                        (issues if selected else unimported).append(finding)
                    if selected and (not issue or issue == "missing"):
                        observations.append(
                            Observation(
                                source_region_code=code,
                                source_region_name=name,
                                value=value,
                                source_cell=coordinate,
                            )
                        )
                        if source.source_id == "institutions_count" and value == 0:
                            issues.append(
                                Issue(
                                    code="zero_denominator_unavailable_for_ratios",
                                    severity="warning",
                                    cell=coordinate,
                                )
                            )
        if not province_rows:
            issues.append(Issue(code="no_provinces", severity="error"))
        if len(codes) != 38:
            issues.append(Issue(code="expected_38_unique_provinces", severity="error"))
        errors = sum(i.severity == "error" for i in issues)
        missing = sum(o.value is None for o in observations)
        return AdditionalStagingResult(
            status="invalid" if errors else "staged",
            source=source,
            dataset_version_id=f"{source.source_id}:sha256:{source.sha256}",
            education_level=education_level,
            observations=tuple(sorted(observations, key=lambda o: o.source_region_code)),
            issues=tuple(issues),
            excluded_aggregates=tuple(aggregates),
            province_rows=province_rows,
            numeric_values=len(observations) - missing,
            missing_values=missing,
            error_count=errors,
            available_levels=tuple(groups),
            unsupported_levels=tuple(level for level in groups if level not in LEVELS),
            unimported_issues=tuple(unimported),
        )
    finally:
        book.close()
