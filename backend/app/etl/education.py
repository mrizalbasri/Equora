"""Stage one education/Total slice without publishing or harmonizing source region codes."""

import hashlib
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Literal
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ConfigDict, Field, HttpUrl, field_validator, model_validator

LEVELS = ("TK", "SD", "SMP", "SMA", "SMK", "SLB")
EducationLevel = Literal["TK", "SD", "SMP", "SMA", "SMK", "SLB"]
ADAPTER_VERSION = "education-total-staging-v1"
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_UNPACKED_BYTES = 20 * 1024 * 1024
MAX_ROWS = 10_000


class FrozenModel(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid")


class SourceMetadata(FrozenModel):
    source_id: str = Field(pattern=r"^[a-z][a-z0-9_-]*$")
    title: str = Field(min_length=1)
    publisher: str = Field(min_length=1)
    source_url: HttpUrl
    sha256: str = Field(pattern=r"^[a-f0-9]{64}$")
    year: int = Field(strict=True, ge=1900, le=9999)
    academic_year: str = Field(pattern=r"^\d{4}/\d{4}$")
    unit: Literal["percent", "school_count"]
    retrieved_at: datetime | None = None
    license_status: str = Field(min_length=1)

    @field_validator("title", "publisher", "license_status")
    @classmethod
    def nonblank(cls, value: str) -> str:
        if not value.strip():
            raise ValueError("Metadata must not be blank")
        return value

    @field_validator("retrieved_at")
    @classmethod
    def timezone_required(cls, value: datetime | None) -> datetime | None:
        if value is not None and value.utcoffset() is None:
            raise ValueError("retrieved_at requires an explicit timezone")
        return value

    @model_validator(mode="after")
    def consistent_period(self):
        start, end = map(int, self.academic_year.split("/"))
        if start != self.year or end != start + 1:
            raise ValueError("Academic year must match year and span consecutive years")
        return self


class Issue(FrozenModel):
    code: str
    severity: Literal["warning", "error"]
    cell: str | None = None


class Observation(FrozenModel):
    source_region_code: str
    source_region_name: str
    value: Decimal | None
    source_cell: str


class StagingResult(FrozenModel):
    adapter_version: str = ADAPTER_VERSION
    status: Literal["staged", "invalid"]
    source: SourceMetadata
    dataset_version_id: str
    education_level: EducationLevel
    source_sheet: Literal["Data"] = "Data"
    ownership: Literal["Total"] = "Total"
    region_namespace: Literal["publisher_source_code_unmapped"] = "publisher_source_code_unmapped"
    observations: tuple[Observation, ...]
    issues: tuple[Issue, ...]
    excluded_aggregates: tuple[str, ...]
    province_rows: int
    numeric_values: int
    missing_values: int
    error_count: int


def _number(value, unit: str) -> tuple[Decimal | None, str | None]:
    if value is None or value in ("", "-", "…"):
        return None, "missing"
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        return None, "invalid_type"
    try:
        number = Decimal(str(value))
    except InvalidOperation:
        return None, "invalid_type"
    if not number.is_finite() or number < 0:
        return None, "invalid_range"
    if unit == "percent" and number > 100:
        return None, "invalid_range"
    if unit == "school_count" and number != number.to_integral_value():
        return None, "invalid_range"
    return number, None


def stage_workbook(
    path: Path, source: SourceMetadata, *, education_level: EducationLevel
) -> StagingResult:
    if education_level not in LEVELS:
        raise ValueError("Select one supported education level explicitly")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("Workbook exceeds staging file size limit")
    payload = path.read_bytes()
    if len(payload) > MAX_FILE_BYTES:
        raise ValueError("Workbook exceeds staging file size limit")
    digest = hashlib.sha256(payload).hexdigest()
    if digest != source.sha256:
        raise ValueError("SHA256 mismatch; review a new source version")
    try:
        with ZipFile(BytesIO(payload)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNPACKED_BYTES:
                raise ValueError("Workbook exceeds unpacked size limit")
    except BadZipFile as exc:
        raise ValueError("Workbook must be an XLSX archive") from exc
    book = load_workbook(BytesIO(payload), read_only=True, data_only=False, keep_links=False)
    try:
        if not {"Data", "Metadata"} <= set(book.sheetnames):
            raise ValueError("Workbook requires Data and Metadata sheets")
        sheet = book["Data"]
        if sheet.max_column != 21 or sheet.max_row > MAX_ROWS:
            raise ValueError("Unexpected workbook dimensions")
        rows = list(sheet.iter_rows())
        header_index = next(
            (
                index
                for index, row in enumerate(rows[:10])
                if tuple(cell.value for cell in row[:3]) == ("No", "Kode Wilayah", "Wilayah")
            ),
            None,
        )
        if header_index is None:
            raise ValueError("Unrecognized source header")
        status_index = next(
            (
                index
                for index in range(header_index + 1, min(header_index + 4, len(rows)))
                if tuple(cell.value for cell in rows[index][3:6]) == ("Negeri", "Swasta", "Total")
            ),
            None,
        )
        if status_index is None:
            raise ValueError("Unrecognized ownership header")
        level_rows = rows[header_index:status_index]
        if not any(
            all(f"({level})" in str(row[3 + index * 3].value) for index, level in enumerate(LEVELS))
            for row in level_rows
        ):
            raise ValueError("Education order changed")
        if any(
            tuple(cell.value for cell in rows[status_index][3 + index * 3 : 6 + index * 3])
            != ("Negeri", "Swasta", "Total")
            for index in range(len(LEVELS))
        ):
            raise ValueError("Ownership order changed")
        observations, issues, aggregates = [], [], []
        codes = set()
        province_rows = missing = numeric = 0
        value_column = 5 + LEVELS.index(education_level) * 3
        for row_number, row in enumerate(rows[status_index + 1 :], status_index + 2):
            if all(cell.value is None for cell in row):
                continue
            name = row[2].value
            if name in ("Indonesia", "Luar Negeri"):
                aggregates.append(name)
                continue
            code = row[1].value
            if (
                not isinstance(name, str)
                or not name.startswith("Provinsi ")
                or type(code) not in (str, int)
                or len(str(code)) != 2
                or not str(code).isascii()
                or not str(code).isdigit()
            ):
                issues.append(Issue(code="invalid_region", severity="error", cell=f"B{row_number}"))
                continue
            province_rows += 1
            code = str(code)
            if code in codes:
                issues.append(
                    Issue(code="duplicate_region", severity="error", cell=f"B{row_number}")
                )
                continue
            codes.add(code)
            cell = row[value_column]
            coordinate = f"{get_column_letter(value_column + 1)}{row_number}"
            value, issue = _number(cell.value, source.unit)
            if cell.data_type == "f":
                issue = "formula_not_allowed"
            if issue:
                issues.append(
                    Issue(
                        code=issue,
                        severity="warning" if issue == "missing" else "error",
                        cell=coordinate,
                    )
                )
                if issue != "missing":
                    continue
                missing += 1
            else:
                numeric += 1
            observations.append(
                Observation(
                    source_region_code=code,
                    source_region_name=name,
                    value=value,
                    source_cell=coordinate,
                )
            )
        if not province_rows:
            issues.append(Issue(code="no_provinces", severity="error"))
        errors = sum(issue.severity == "error" for issue in issues)
        return StagingResult(
            status="invalid" if errors else "staged",
            source=source,
            dataset_version_id=f"{source.source_id}:sha256:{digest}",
            education_level=education_level,
            observations=tuple(sorted(observations, key=lambda row: row.source_region_code)),
            issues=tuple(issues),
            excluded_aggregates=tuple(aggregates),
            province_rows=province_rows,
            numeric_values=numeric,
            missing_values=missing,
            error_count=errors,
        )
    finally:
        book.close()
