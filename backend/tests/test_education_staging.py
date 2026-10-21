"""Synthetic workbooks; no official observations are checked into Git."""

import hashlib
import json
from datetime import UTC, datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook
from pydantic import ValidationError

from app.etl import education, stage_education
from app.etl.education import LEVELS, SourceMetadata, stage_workbook


def workbook(tmp_path, values=(0, None, 75), *, unit="percent", change=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    book.create_sheet("Metadata")
    sheet.append(["Synthetic source, not province statistics"])
    sheet.append(["No", "Kode Wilayah", "Wilayah"])
    for index, level in enumerate(LEVELS):
        column = 4 + index * 3
        sheet.cell(3, column, f"Synthetic ({level})")
        for offset, ownership in enumerate(("Negeri", "Swasta", "Total")):
            sheet.cell(4, column + offset, ownership)
    for index, value in enumerate(values):
        row = 5 + index
        sheet.cell(row, 1, index + 1)
        sheet.cell(row, 2, str(10 + index))
        sheet.cell(row, 3, f"Provinsi Sintetis {index}")
        for column in range(4, 22):
            sheet.cell(row, column, value)
    sheet.append([None, "00", "Indonesia"] + [99] * 18)
    sheet.append([None, "99", "Luar Negeri"] + [99] * 18)
    if change:
        change(book)
    path = tmp_path / "synthetic.xlsx"
    book.save(path)
    book.close()
    source = SourceMetadata(
        source_id="synthetic",
        title="Synthetic source",
        publisher="Test fixture",
        source_url="https://example.org/synthetic",
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        year=2025,
        academic_year="2025/2026",
        unit=unit,
        retrieved_at=datetime(2026, 10, 21, tzinfo=UTC),
        license_status="synthetic_fixture_only",
    )
    return path, source


def test_golden_staging_keeps_zero_missing_and_provenance(tmp_path):
    path, source = workbook(tmp_path)
    result = stage_workbook(path, source, education_level="SD")
    assert result.status == "staged"
    assert [row.value for row in result.observations] == [Decimal(0), None, Decimal(75)]
    assert [row.source_cell for row in result.observations] == ["I5", "I6", "I7"]
    assert result.province_rows == 3
    assert result.numeric_values == 2
    assert result.missing_values == 1
    assert result.error_count == 0
    assert result.excluded_aggregates == ("Indonesia", "Luar Negeri")
    assert result.source == source
    assert result.dataset_version_id == f"synthetic:sha256:{source.sha256}"
    assert result.region_namespace == "publisher_source_code_unmapped"
    assert result.issues[0].code == "missing"
    assert result.issues[0].cell == "I6"
    serialized = json.loads(result.model_dump_json())
    assert serialized["observations"][0]["value"] == "0"
    assert serialized["observations"][1]["value"] is None
    assert stage_workbook(path, source, education_level="SD") == result


@pytest.mark.parametrize("token", [None, "", "-", "…"])
def test_missing_tokens_stay_null(tmp_path, token):
    path, source = workbook(tmp_path, (token,))
    result = stage_workbook(path, source, education_level="SD")
    assert result.status == "staged"
    assert result.observations[0].value is None
    assert result.missing_values == 1


@pytest.mark.parametrize(
    ("value", "unit", "issue"),
    [
        (True, "percent", "invalid_type"),
        ("75", "percent", "invalid_type"),
        (-1, "percent", "invalid_range"),
        (101, "percent", "invalid_range"),
        (1.5, "school_count", "invalid_range"),
        ("=1+2", "percent", "formula_not_allowed"),
    ],
)
def test_invalid_values_are_errors_not_missing(tmp_path, value, unit, issue):
    path, source = workbook(tmp_path, (value,), unit=unit)
    result = stage_workbook(path, source, education_level="SD")
    assert result.status == "invalid"
    assert result.error_count == 1
    assert result.missing_values == 0
    assert result.observations == ()
    assert result.issues[0].code == issue


@pytest.mark.parametrize("value", [float("inf"), float("nan"), Decimal("NaN")])
def test_nonfinite_values_are_rejected(value):
    assert education._number(value, "percent") == (None, "invalid_range")


def test_counts_are_not_converted_to_percent(tmp_path):
    path, source = workbook(tmp_path, (1250,), unit="school_count")
    result = stage_workbook(path, source, education_level="TK")
    assert result.status == "staged"
    assert result.observations[0].value == Decimal(1250)
    assert result.source.unit == "school_count"
    assert result.observations[0].source_cell == "F5"


def test_only_selected_education_total_is_imported(tmp_path):
    def change(book):
        book["Data"]["G5"] = 50
        book["Data"]["H5"] = 80
        book["Data"]["I5"] = 65
        book["Data"]["L5"] = 90

    path, source = workbook(tmp_path, (0,), change=change)
    assert stage_workbook(path, source, education_level="SD").observations[0].value == 65
    assert stage_workbook(path, source, education_level="SMP").observations[0].value == 90
    with pytest.raises(ValueError, match="supported education"):
        stage_workbook(path, source, education_level="all")


@pytest.mark.parametrize(
    ("coordinate", "value", "message"),
    [
        ("A2", "changed", "source header"),
        ("G3", "SMP", "Education order"),
        ("I4", "Negeri", "Ownership order"),
    ],
)
def test_layout_changes_require_adapter_review(tmp_path, coordinate, value, message):
    path, source = workbook(
        tmp_path, change=lambda book: setattr(book["Data"][coordinate], "value", value)
    )
    with pytest.raises(ValueError, match=message):
        stage_workbook(path, source, education_level="SD")


def test_changed_file_is_rejected_before_parsing(tmp_path):
    path, source = workbook(tmp_path)
    path.write_bytes(b"changed file")
    with pytest.raises(ValueError, match="SHA256 mismatch"):
        stage_workbook(path, source, education_level="SD")


def test_non_xlsx_payload_is_rejected_even_with_matching_hash(tmp_path):
    path, source = workbook(tmp_path)
    path.write_bytes(b"not an xlsx archive")
    source = source.model_copy(update={"sha256": hashlib.sha256(path.read_bytes()).hexdigest()})
    with pytest.raises(ValueError, match="XLSX archive"):
        stage_workbook(path, source, education_level="SD")


def test_missing_metadata_sheet_is_rejected(tmp_path):
    path, source = workbook(tmp_path, change=lambda book: book.remove(book["Metadata"]))
    with pytest.raises(ValueError, match="Data and Metadata"):
        stage_workbook(path, source, education_level="SD")


@pytest.mark.parametrize("academic_year", ["2024/2025", "2025/2027"])
def test_conflicting_metadata_period_is_rejected(tmp_path, academic_year):
    _, source = workbook(tmp_path)
    data = source.model_dump()
    data["academic_year"] = academic_year
    with pytest.raises(ValidationError, match="Academic year"):
        SourceMetadata(**data)


def test_duplicate_source_region_is_not_silently_overwritten(tmp_path):
    path, source = workbook(
        tmp_path, change=lambda book: setattr(book["Data"]["B6"], "value", "10")
    )
    result = stage_workbook(path, source, education_level="SD")
    assert result.status == "invalid"
    assert result.error_count == 1
    assert result.issues[0].code == "duplicate_region"
    assert len({row.source_region_code for row in result.observations}) == len(result.observations)


@pytest.mark.parametrize("code", [True, "1", "123", "ab", "１２"])
def test_unrecognized_region_codes_fail_closed(tmp_path, code):
    path, source = workbook(
        tmp_path, (1,), change=lambda book: setattr(book["Data"]["B5"], "value", code)
    )
    result = stage_workbook(path, source, education_level="SD")
    assert result.status == "invalid"
    assert any(issue.code == "invalid_region" for issue in result.issues)


def test_no_province_rows_is_invalid(tmp_path):
    path, source = workbook(tmp_path, ())
    result = stage_workbook(path, source, education_level="SD")
    assert result.status == "invalid"
    assert result.issues[0].code == "no_provinces"


@pytest.mark.parametrize("limit", ["MAX_FILE_BYTES", "MAX_UNPACKED_BYTES", "MAX_ROWS"])
def test_staging_resource_limits(tmp_path, monkeypatch, limit):
    path, source = workbook(tmp_path)
    monkeypatch.setattr(education, limit, 1)
    with pytest.raises(ValueError):
        stage_workbook(path, source, education_level="SD")


def test_metadata_requires_timezone_and_explicit_terms_state(tmp_path):
    _, source = workbook(tmp_path)
    data = source.model_dump()
    data["retrieved_at"] = datetime(2026, 10, 21)
    with pytest.raises(ValidationError, match="timezone"):
        SourceMetadata(**data)
    data["retrieved_at"] = source.retrieved_at
    data["license_status"] = " "
    with pytest.raises(ValidationError, match="blank"):
        SourceMetadata(**data)


def test_output_is_exclusive_and_confined_to_processed(tmp_path, monkeypatch):
    path, source = workbook(tmp_path)
    result = stage_workbook(path, source, education_level="SD")
    monkeypatch.setattr(stage_education, "PROCESSED_ROOT", tmp_path / "processed")
    output = tmp_path / "processed/staging/synthetic.json"
    stage_education.write_staging(output, result)
    original = output.read_bytes()
    with pytest.raises(FileExistsError):
        stage_education.write_staging(output, result)
    assert output.read_bytes() == original
    with pytest.raises(ValueError, match="data/processed"):
        stage_education.write_staging(tmp_path / "outside.json", result)
    assert not (tmp_path / "outside.json").exists()


def test_cli_uses_audited_metadata_and_reports_staging_only(tmp_path, monkeypatch, capsys):
    path, source = workbook(tmp_path)
    manifest = {
        "publisher": source.publisher,
        "academic_year": source.academic_year,
        "license_status": source.license_status,
        "sources": [
            {
                "id": source.source_id,
                "sha256": source.sha256,
                "source_url": str(source.source_url),
                "unit": source.unit,
            }
        ],
    }
    audit = {
        "sources": [
            {
                "id": source.source_id,
                "sha256": source.sha256,
                "title": source.title,
                "retrieved_at": source.retrieved_at.isoformat(),
            }
        ]
    }
    metadata = tmp_path / "data/metadata"
    metadata.mkdir(parents=True)
    (metadata / "source-candidates.json").write_text(json.dumps(manifest))
    (metadata / "source-audit.json").write_text(json.dumps(audit))
    monkeypatch.setattr(stage_education, "ROOT", tmp_path)
    monkeypatch.setattr(stage_education, "PROCESSED_ROOT", tmp_path / "data/processed")
    output = tmp_path / "data/processed/staging.json"
    arguments = [
        "--source-id",
        "synthetic",
        "--file",
        str(path),
        "--education-level",
        "SD",
        "--output",
        str(output),
    ]
    assert stage_education.main(arguments) == 0
    assert json.loads(capsys.readouterr().out) == {
        "status": "staged",
        "observations": 3,
        "missing": 1,
        "errors": 0,
        "release_status": "not_approved",
    }
    assert json.loads(output.read_text())["status"] == "staged"
    with pytest.raises(SystemExit) as exc:
        stage_education.main(arguments)
    assert exc.value.code == 2
    audit["sources"][0]["sha256"] = "0" * 64
    with pytest.raises(ValueError, match="do not match"):
        stage_education.audited_source("synthetic", manifest, audit)


def test_cli_preserves_invalid_report_and_returns_failure(tmp_path, monkeypatch, capsys):
    path, source = workbook(tmp_path, (0, 101))
    monkeypatch.setattr(stage_education, "PROCESSED_ROOT", tmp_path / "processed")
    monkeypatch.setattr(stage_education, "audited_source", lambda *args: source)
    output = tmp_path / "processed/invalid.json"
    assert (
        stage_education.main(
            [
                "--source-id",
                "synthetic",
                "--file",
                str(path),
                "--education-level",
                "SD",
                "--output",
                str(output),
            ]
        )
        == 1
    )
    assert json.loads(capsys.readouterr().out)["errors"] == 1
    report = json.loads(output.read_text())
    assert report["status"] == "invalid"
    assert report["issues"][0]["cell"] == "I6"
    assert len(report["observations"]) == 1


def test_unknown_manifest_source_is_rejected():
    with pytest.raises(ValueError, match="both the manifest and audit"):
        stage_education.audited_source("unknown", {"sources": []}, {"sources": []})
