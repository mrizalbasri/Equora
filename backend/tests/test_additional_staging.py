"""Synthetic contracts for the six additions, separate from the four legacy sources."""

import hashlib
import json
from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook

from app.etl import additional_education as adapter
from app.etl import register_education, stage_education
from app.etl.education import LEVELS
from app.etl.local_release import read_source
from app.etl.registry import register_staging


def fixture(tmp_path, source_id="internet_percent", *, value=None, change=None):
    unit, scale, layout = adapter.PROFILES[source_id]
    title = "Synthetic fixture, Tahun Ajaran 2025/2026"
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    metadata = book.create_sheet("Metadata")
    metadata.append(["Nama", title])
    metadata.append(["Satuan", "Persen" if unit == "percent" else "Unit"])
    sheet.append([title])
    sheet.append(["No", "Kode Wilayah", "Wilayah"])
    levels = adapter.DENOMINATOR_LEVELS if source_id == "institutions_count" else LEVELS
    for index, level in enumerate(levels):
        start = 4 + index * 3
        sheet.cell(3, start, f"Synthetic ({level})")
        for offset, owner in enumerate(("Negeri", "Swasta", "Total")):
            sheet.cell(4, start + offset, owner)
    if source_id == "institutions_count":
        sheet.cell(1, 53).number_format = "0"
    for index in range(38):
        row = index + 5
        sheet.cell(row, 1, index + 1)
        sheet.cell(row, 2, str(index + 10))
        sheet.cell(row, 3, f"Provinsi Synthetic {index}")
        for column in range(4, 4 + 3 * len(levels)):
            cell = sheet.cell(
                row, column, value if value is not None else (0.0075 if unit == "percent" else 1250)
            )
            if unit == "percent":
                cell.number_format = "0.00%"
    sheet.append([None, "00", "Indonesia"])
    sheet.append([None, "99", "Luar Negeri"])
    if change:
        change(book)
    path = tmp_path / "synthetic.xlsx"
    book.save(path)
    book.close()
    record = {
        "id": source_id,
        "filename": path.name,
        "title": title,
        "source_url": "https://example.org/synthetic",
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "unit": unit,
        "storage_scale": scale,
        "downloaded_at": None,
    }
    manifest = {"academic_year": "2025/2026", "sources": [record]}
    report = {
        **record,
        "education_levels": list(levels),
        "columns_in_file": 53 if source_id == "institutions_count" else 21,
        "province_rows": 38,
    }
    audit = {"sources": [report]}
    source = adapter.audited_additional_source(source_id, manifest, audit)
    return path, source, manifest, audit


@pytest.mark.parametrize("source_id", adapter.PROFILES)
def test_six_profiles_keep_units_scale_and_provenance(tmp_path, source_id):
    path, source, _, _ = fixture(tmp_path, source_id)
    result = adapter.stage_additional_workbook(path, source, education_level="SD")
    assert result.status == "staged"
    assert result.numeric_values == 38
    assert result.source.retrieved_at is None
    assert result.observations[0].value == (Decimal("0.75") if source.unit == "percent" else 1250)
    assert result.observations[0].source_cell == (
        "AA5" if source_id == "institutions_count" else "I5"
    )
    assert result.excluded_aggregates == ("Indonesia", "Luar Negeri")
    assert "retrieved_at_unknown" in {issue.code for issue in result.issues}
    assert json.loads(result.model_dump_json())["source"]["storage_scale"] == source.storage_scale
    assert adapter.stage_additional_workbook(path, source, education_level="SD") == result


@pytest.mark.parametrize(
    "value,code",
    [
        (0, None),
        (1, None),
        (None, "missing"),
        ("-", "missing"),
        (1.01, "invalid_range"),
        (-1, "invalid_range"),
        (True, "invalid_type"),
        ("0.75", "invalid_type"),
        ("=3/4", "formula_not_allowed"),
    ],
)
def test_fraction_values_are_not_guessed_or_filled(tmp_path, value, code):
    path, source, _, _ = fixture(
        tmp_path, change=lambda book: setattr(book["Data"]["I5"], "value", value)
    )
    result = adapter.stage_additional_workbook(path, source, education_level="SD")
    if code:
        assert any(issue.code == code and issue.cell == "I5" for issue in result.issues)
        assert result.status == ("staged" if code == "missing" else "invalid")
        assert result.missing_values == (1 if code == "missing" else 0)
    else:
        assert result.observations[0].value == value * 100


@pytest.mark.parametrize(
    "source_id,format",
    [("internet_percent", "General"), ("computers_percent", "0%"), ("water_count", "0.00%")],
)
def test_format_changes_require_review(tmp_path, source_id, format):
    path, source, _, _ = fixture(
        tmp_path,
        source_id,
        change=lambda book: setattr(book["Data"]["I5"], "number_format", format),
    )
    result = adapter.stage_additional_workbook(path, source, education_level="SD")
    assert result.status == "invalid"
    assert any(
        issue.code == "storage_format_changed" and issue.cell == "I5" for issue in result.issues
    )


def test_denominator_formula_and_missing_ownership_remain_visible_outside_slice(tmp_path):
    def change(book):
        book["Data"]["I5"] = "=1+2"
        book["Data"]["Z5"] = None
        book["Data"]["AA5"] = 0

    path, source, _, _ = fixture(tmp_path, "institutions_count", change=change)
    result = adapter.stage_additional_workbook(path, source, education_level="SD")
    assert result.status == "staged"
    assert result.observations[0].value == 0
    assert {(i.code, i.cell) for i in result.unimported_issues} == {
        ("formula_not_allowed", "I5"),
        ("missing", "Z5"),
    }
    assert "KB" in result.unsupported_levels
    assert result.unimported_ownership == ("Negeri", "Swasta")
    assert any(i.code == "zero_denominator_unavailable_for_ratios" for i in result.issues)
    with pytest.raises(ValueError, match="PAUD/Dikmas"):
        adapter.stage_additional_workbook(path, source, education_level="KB")


@pytest.mark.parametrize(
    "coordinate,value,message",
    [
        ("D3", "Unknown (XX)", "header groups"),
        ("Y3", "Duplicate (TK)", "header groups"),
        ("Y3", None, "header groups"),
        ("AA4", "Swasta", "Ownership order"),
        ("BA5", 1, "outside recognized"),
        ("A1", "changed title", "title/metadata"),
    ],
)
def test_denominator_layout_and_unmapped_data_fail_closed(tmp_path, coordinate, value, message):
    path, source, _, _ = fixture(
        tmp_path,
        "institutions_count",
        change=lambda book: setattr(book["Data"][coordinate], "value", value),
    )
    with pytest.raises(ValueError, match=message):
        adapter.stage_additional_workbook(path, source, education_level="SD")


@pytest.mark.parametrize(
    "change,code",
    [
        (lambda book: setattr(book["Data"]["B6"], "value", "10"), "duplicate_region"),
        (lambda book: book["Data"].delete_rows(5), "expected_38_unique_provinces"),
        (lambda book: setattr(book["Data"]["AA5"], "value", "=1+2"), "formula_not_allowed"),
        (lambda book: setattr(book["Data"]["AA5"], "value", 1.5), "invalid_range"),
    ],
)
def test_selected_errors_and_incomplete_coverage_are_reported(tmp_path, change, code):
    path, source, _, _ = fixture(tmp_path, "institutions_count", change=change)
    result = adapter.stage_additional_workbook(path, source, education_level="SD")
    assert result.status == "invalid"
    assert any(issue.code == code for issue in result.issues)


@pytest.mark.parametrize(
    "key,value",
    [
        ("sha256", "0" * 64),
        ("storage_scale", "percentage_points"),
        ("education_levels", ["SD"]),
        ("province_rows", 37),
    ],
)
def test_manifest_and_audit_must_agree(tmp_path, key, value):
    _, _, manifest, audit = fixture(tmp_path)
    audit["sources"][0][key] = value
    with pytest.raises(ValueError):
        adapter.audited_additional_source("internet_percent", manifest, audit)


def test_source_profile_cannot_be_used_with_a_different_scale(tmp_path):
    path, source, manifest, audit = fixture(tmp_path)
    changed = source.model_copy(update={"storage_scale": "percentage_points"})
    with pytest.raises(ValueError, match="storage_scale|storage scale"):
        adapter.stage_additional_workbook(path, changed, education_level="SD")
    with pytest.raises(ValueError, match="original adapter"):
        adapter.audited_additional_source("water_percent", manifest, audit)


def test_additional_hash_and_resource_guards(tmp_path, monkeypatch):
    path, source, _, _ = fixture(tmp_path)
    monkeypatch.setattr(adapter, "MAX_ROWS", 1)
    with pytest.raises(ValueError, match="dimensions"):
        adapter.stage_additional_workbook(path, source, education_level="SD")
    path.write_bytes(b"not a workbook")
    with pytest.raises(ValueError, match="SHA256"):
        adapter.stage_additional_workbook(path, source, education_level="SD")


def write_metadata(tmp_path, manifest, audit):
    metadata = tmp_path / "data/metadata"
    metadata.mkdir(parents=True)
    for filename, payload in (
        ("additional-source-candidates.json", manifest),
        ("additional-source-audit.json", audit),
    ):
        (metadata / filename).write_text(json.dumps(payload))


def test_additional_staging_cli_writes_explicit_report(tmp_path, monkeypatch, capsys):
    path, _, manifest, audit = fixture(
        tmp_path,
        "institutions_count",
        change=lambda book: setattr(book["Data"]["I5"], "value", "=1+2"),
    )
    write_metadata(tmp_path, manifest, audit)
    monkeypatch.setattr(stage_education, "ROOT", tmp_path)
    monkeypatch.setattr(stage_education, "PROCESSED_ROOT", tmp_path / "data/processed")
    output = tmp_path / "data/processed/staging.json"
    args = [
        "--source-set",
        "additional",
        "--source-id",
        "institutions_count",
        "--file",
        str(path),
        "--education-level",
        "SD",
        "--output",
        str(output),
    ]
    assert stage_education.main(args) == 0
    summary = json.loads(capsys.readouterr().out)
    assert summary["unimported_issues"] == 1
    report = json.loads(output.read_text())
    assert report["unimported_issues"][0]["cell"] == "I5"
    assert report["source"]["retrieved_at"] is None
    assert report["adapter_version"] == adapter.ADAPTER_VERSION


def test_additional_import_cli_uses_new_adapter_and_retains_draft_status(
    tmp_path, monkeypatch, capsys
):
    path, _, manifest, audit = fixture(tmp_path)
    write_metadata(tmp_path, manifest, audit)
    monkeypatch.setattr(register_education, "ROOT", tmp_path)
    monkeypatch.setattr(
        register_education,
        "get_settings",
        lambda: SimpleNamespace(database_url="postgresql://localhost/equora"),
    )
    engine = MagicMock()
    monkeypatch.setattr(register_education, "build_engine", lambda settings: engine)
    monkeypatch.setattr(register_education, "Session", MagicMock())
    seen = []

    def register(session, result):
        seen.append(result)
        return SimpleNamespace(import_id="synthetic", created=True, status=result.status)

    monkeypatch.setattr(register_education, "register_staging", register)
    assert (
        register_education.main(
            [
                "--source-set",
                "additional",
                "--source-id",
                "internet_percent",
                "--file",
                str(path),
                "--education-level",
                "SD",
            ]
        )
        == 0
    )
    assert seen[0].observations[0].value == Decimal("0.75")
    assert seen[0].source.storage_scale == "fraction_with_percent_format"
    assert json.loads(capsys.readouterr().out)["release_status"] == "not_approved"
    engine.dispose.assert_called_once()


def test_local_review_uses_same_adapter_and_rejects_invalid_totals(tmp_path):
    path, source, manifest, _ = fixture(
        tmp_path,
        "institutions_count",
        change=lambda book: setattr(book["Data"]["I5"], "value", "=1+2"),
    )
    results = read_source(path, manifest["sources"][0])
    assert results["SD"].adapter_version == adapter.ADAPTER_VERSION
    assert results["SD"].source == source
    assert results["SD"].unimported_issues[0].code == "formula_not_allowed"
    path, _, manifest, _ = fixture(tmp_path, value=2)
    with pytest.raises(ValueError, match="stage it for a detailed report"):
        read_source(path, manifest["sources"][0])


def test_registry_serializes_additional_source_contract_and_unimported_findings(tmp_path):
    path, source, _, _ = fixture(
        tmp_path,
        "institutions_count",
        change=lambda book: setattr(book["Data"]["I5"], "value", "=1+2"),
    )
    result = adapter.stage_additional_workbook(path, source, education_level="SD")
    session = MagicMock()
    session.get.return_value = SimpleNamespace(source_metadata=source.model_dump(mode="json"))
    registration = register_staging(session, result)
    assert registration.created
    version_params = session.execute.call_args_list[1].args[0].compile().params
    assert version_params["source_metadata"]["storage_scale"] == "count"
    assert version_params["source_metadata"]["retrieved_at"] is None
    import_params = session.execute.call_args_list[2].args[0].compile().params
    assert import_params["summary"]["unsupported_levels"] == list(result.unsupported_levels)
    assert import_params["summary"]["unimported_issues"] == [
        {"code": "formula_not_allowed", "severity": "error", "cell": "I5"}
    ]
    assert import_params["status"] == "staged"


def test_additional_cli_keeps_invalid_report_without_filling_value(tmp_path, monkeypatch, capsys):
    path, _, manifest, audit = fixture(tmp_path, value=2)
    write_metadata(tmp_path, manifest, audit)
    monkeypatch.setattr(stage_education, "ROOT", tmp_path)
    monkeypatch.setattr(stage_education, "PROCESSED_ROOT", tmp_path / "data/processed")
    output = tmp_path / "data/processed/invalid.json"
    assert (
        stage_education.main(
            [
                "--source-set",
                "additional",
                "--source-id",
                "internet_percent",
                "--file",
                str(path),
                "--education-level",
                "SD",
                "--output",
                str(output),
            ]
        )
        == 1
    )
    assert json.loads(capsys.readouterr().out)["errors"] == 38
    report = json.loads(output.read_text())
    assert report["status"] == "invalid"
    assert report["observations"] == []
