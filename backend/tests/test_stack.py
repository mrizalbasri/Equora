import os
from pathlib import Path

import polars as pl
import pytest
from alembic import command
from alembic.config import Config
from openpyxl import Workbook, load_workbook
from ortools.sat.python import cp_model
from sqlalchemy import inspect
from sqlalchemy.engine import make_url

from app.core.config import Settings
from app.db.session import build_engine


def test_data_and_solver_dependencies(tmp_path: Path):
    workbook = Workbook()
    workbook.active.append(["value"])
    workbook.active.append([3])
    path = tmp_path / "synthetic.xlsx"
    workbook.save(path)
    loaded = load_workbook(path, read_only=True)
    try:
        frame = pl.DataFrame({"value": [loaded.active["A2"].value]})
        assert frame["value"].sum() == 3
    finally:
        loaded.close()

    model = cp_model.CpModel()
    packages = model.new_int_var(0, 5, "synthetic_packages")
    model.add(3 * packages <= 10)
    model.maximize(packages)
    solver = cp_model.CpSolver()
    assert solver.solve(model) == cp_model.OPTIMAL
    assert solver.value(packages) == 3


@pytest.mark.skipif(not os.getenv("TEST_DATABASE_URL"), reason="TEST_DATABASE_URL is not set")
def test_postgres_migration_round_trip(monkeypatch):
    url = os.environ["TEST_DATABASE_URL"]
    if make_url(url).database != "equora_test":
        pytest.fail("Migration round-trip requires the dedicated equora_test database")
    monkeypatch.setenv("DATABASE_URL", url)
    from app.core.config import get_settings

    get_settings.cache_clear()
    config = Config(str(Path(__file__).resolve().parents[1] / "alembic.ini"))
    engine = build_engine(Settings(database_url=url, _env_file=None))
    try:
        command.upgrade(config, "head")
        assert "alembic_version" in inspect(engine).get_table_names()
        assert "staged_observations" in inspect(engine).get_table_names()
        command.downgrade(config, "base")
        command.upgrade(config, "head")
    finally:
        engine.dispose()
        get_settings.cache_clear()
