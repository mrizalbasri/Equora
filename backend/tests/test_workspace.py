"""Synthetic end-to-end decision and transaction cases; never read user workbooks."""

import copy
import hashlib
import os
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pydantic import ValidationError

from app.api.workspace import ScenarioInput, session_for
from app.core.config import Settings
from app.db.registry import LocalRelease
from app.etl.education import LEVELS
from app.etl.local_release import CORE, read_source
from app.main import create_app
from app.services.workspace import calculate_scenario, compare_snapshots, priority


@pytest.fixture
def payload():
    sources = [
        {
            "source_id": key,
            "unit": "percent",
            "source_url": "https://example.org/synthetic",
            "title": "Synthetic " + key,
            "sha256": "a" * 64,
            "license_status": "synthetic_fixture_only",
            "quality": {"SD": {"missing": 0, "numeric": 4}},
            "year": 2025,
            "academic_year": "2025/2026",
            "publisher": "Synthetic",
            "definition": "Synthetic tests only",
            "storage_scale": "percentage_points",
        }
        for key in CORE
    ]
    rows = [
        {
            "code": code,
            "name": "Synthetic " + code,
            "values": {**dict.fromkeys(CORE, value), "institutions_count": denominator},
            "cells": dict.fromkeys(CORE, "I5"),
        }
        for code, value, denominator in (
            ("11", "20", "10"),
            ("12", "60", "20"),
            ("13", "0", "0"),
            ("14", None, "10"),
        )
    ]
    return {
        "id": "synthetic-release-workspace",
        "status": "local_review",
        "year": 2025,
        "ownership": "Total",
        "adapter_version": "synthetic",
        "region_namespace": "synthetic",
        "academic_year": "2025/2026",
        "sources": sources,
        "scopes": {"SD": rows},
        "limitations": ["Synthetic tests only"],
    }


def params(payload, **kwargs):
    return ScenarioInput(
        release_id=payload["id"],
        name="Synthetic <script>",
        weights=dict.fromkeys(CORE, Decimal("0.25")),
        budget_idr=30,
        cost_idr=10,
        capacity=2,
        **kwargs,
    )


def test_needs_denominator_missing_and_zero_weight(payload):
    result = priority(payload, "SD", dict.fromkeys(CORE, Decimal("0.25")))
    assert [(row.region_code, row.score) for row in result.ranked] == [
        ("11", Decimal("0.8")),
        ("12", Decimal("0.4")),
    ]
    assert set(result.excluded) == {"13", "14"}
    payload["scopes"]["SD"][3]["values"].update(dict.fromkeys(CORE, "30"))
    payload["scopes"]["SD"][3]["values"][CORE[0]] = None
    weights = dict.fromkeys(CORE, Decimal(0))
    weights[CORE[1]] = Decimal(1)
    assert "14" in {row.region_code for row in priority(payload, "SD", weights).ranked}


def test_scenario_budget_baseline_and_compare(payload):
    result = calculate_scenario(payload, params(payload))
    assert result["allocation"]["status"] == "optimal"
    assert result["allocation"]["allocations"] == {"11": 2, "12": 1}
    assert result["allocation"]["evaluation"]["budget_used_idr"] == 30
    assert result["region_names"]["11"] == "Synthetic 11"
    for baseline in result["baselines"].values():
        assert baseline["evaluation"]["budget_used_idr"] <= 30
    other = copy.deepcopy(result)
    other["parameters"]["weights"][CORE[0]] = "0.50"
    assert compare_snapshots(result, other)["objective_comparable"] is False
    other["release_id"] = "other"
    with pytest.raises(ValueError, match="Snapshot berbeda"):
        compare_snapshots(result, other)


def test_invalid_assumptions_and_infeasible(payload):
    with pytest.raises(ValidationError):
        params(payload, region_codes=[])
    bad = params(payload).model_copy(update={"minimum_packages": 3})
    with pytest.raises(ValueError, match="minimum_packages"):
        calculate_scenario(payload, bad)
    infeasible = params(payload).model_copy(update={"budget_idr": 0, "minimum_packages": 1})
    result = calculate_scenario(payload, infeasible)
    assert result["allocation"]["status"] == "infeasible"
    assert result["allocation"]["allocations"] is None
    with pytest.raises(ValueError, match="solusi"):
        compare_snapshots(result, result)


@pytest.mark.parametrize(
    "wide,scale,value,expected",
    [
        (False, "fraction_with_percent_format", 0.625, Decimal("62.500")),
        (False, "percentage_points", 62.5, Decimal("62.5")),
        (True, "count", 12, Decimal(12)),
    ],
)
def test_explicit_headers_and_percentage_scale(tmp_path, wide, scale, value, expected):
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["Synthetic source"])
    sheet.append(["No", "Kode Wilayah", "Wilayah"])
    offset = 21 if wide else 3
    for index, level in enumerate(LEVELS):
        column = offset + index * 3 + 1
        sheet.cell(2, column, f"Synthetic ({level})")
        for delta, owner in enumerate(("Negeri", "Swasta", "Total")):
            sheet.cell(3, column + delta, owner)
    # The denominator's first group precedes the six selected education groups.
    for delta, owner in enumerate(("Negeri", "Swasta", "Total")):
        sheet.cell(3, 4 + delta, owner)
    for index in range(38):
        row = index + 4
        sheet.cell(row, 1, index + 1)
        sheet.cell(row, 2, str(index + 10))
        sheet.cell(row, 3, f"Provinsi Synthetic {index}")
        for j in range(6):
            cell = sheet.cell(row, offset + j * 3 + 3, value)
            if scale == "fraction_with_percent_format":
                cell.number_format = "0.00%"
    path = tmp_path / "source.xlsx"
    book.save(path)
    record = {
        "id": "synthetic",
        "title": "Synthetic source",
        "source_url": "https://example.org",
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "unit": "school_count" if wide else "percent",
        "storage_scale": scale,
    }
    result = read_source(path, record)
    assert all(item.value == expected for item in result["SD"].observations)
    assert result["SD"].observations[0].source_cell == ("AA4" if wide else "I4")
    path.write_bytes(path.read_bytes() + b"changed")
    with pytest.raises(ValueError, match="SHA256"):
        read_source(path, record)


def test_write_guard_and_admin_auth():
    with TestClient(create_app(Settings(_env_file=None))) as client:
        assert client.post("/api/v1/admin/import", json={}).status_code == 401
        assert client.post("/api/v1/scenarios", content="x" * 65537).status_code == 413
        assert (
            client.post(
                "/api/v1/scenarios", json={}, headers={"Origin": "https://evil.example"}
            ).status_code
            == 403
        )


def test_public_profiles_preserve_requested_release(payload, monkeypatch):
    import importlib

    module = importlib.import_module("app.api.workspace")
    older = copy.deepcopy(payload)
    older["id"] = "synthetic-older"
    older["scopes"]["SD"][0]["values"][CORE[0]] = "12"

    def release_for(_session, release_id=None):
        if release_id == older["id"]:
            return older
        if release_id in (None, payload["id"]):
            return payload
        from fastapi import HTTPException

        raise HTTPException(404, "Rilis tidak ditemukan")

    monkeypatch.setattr(module, "local_release", release_for)
    app = create_app(Settings(_env_file=None))
    app.dependency_overrides[session_for] = lambda: None
    with TestClient(app) as client:
        query = {"release_id": older["id"], "education_level": "SD"}
        assert client.get("/api/v1/workspace", params=query).json()["id"] == older["id"]
        assert client.get("/api/v1/regions", params=query).json()[0]["code"] == "11"
        result = client.get("/api/v1/regions/11", params=query).json()
        assert result["release_id"] == older["id"]
        assert result["values"][CORE[0]] == "12"
        assert result["status"] == "local_review"
        assert result["academic_year"] == "2025/2026"
        assert result["ownership"] == "Total"
        for path in ("workspace", "regions", "regions/11"):
            assert client.get("/api/v1/" + path, params={"release_id": "absent"}).status_code == 404
        assert client.get("/api/v1/regions/99", params=query).status_code == 404
        for path in ("regions", "regions/11"):
            response = client.get("/api/v1/" + path, params={"education_level": "SMP"})
            assert response.status_code == 422


@pytest.mark.skipif(not os.getenv("TEST_DATABASE_URL"), reason="TEST_DATABASE_URL is not set")
def test_postgres_snapshot_routes(payload, session):
    session.add(LocalRelease(id=payload["id"], payload=payload))
    session.flush()
    app = create_app(Settings(_env_file=None))
    app.dependency_overrides[session_for] = lambda: session
    with TestClient(app) as client:
        assert client.get("/api/v1/workspace").status_code == 200
        response = client.post("/api/v1/scenarios", json=params(payload).model_dump(mode="json"))
        assert response.status_code == 201, response.text
        saved = response.json()
        route = "/api/v1/scenarios/" + saved["id"]
        assert client.get(route).json() == saved
        assert client.get(route + "/export").text.find("synthetic_fixture_only") > 0
        assert client.get(route + "/pdf").content.startswith(b"%PDF")
        compared_pdf = client.post(
            "/api/v1/scenarios/compare/pdf", json={"left": saved["id"], "right": saved["id"]}
        )
        assert compared_pdf.status_code == 200
        assert compared_pdf.content.startswith(b"%PDF")
        assert client.patch(route, json={"name": "change"}).status_code == 405
        assert (
            client.post(
                "/api/v1/scenarios/compare",
                json={"left": saved["id"], "right": saved["id"]},
            ).status_code
            == 200
        )
        observed = client.get("/api/v1/observations").json()
        assert observed["rows"][2]["status"] == "no_schools"
        assert observed["rows"][2]["source_value"] == "0"
        assert observed["rows"][2]["value"] is None
        assert client.get("/api/v1/datasets/absent").status_code == 404
        older = copy.deepcopy(payload)
        older["id"] = "synthetic-profile-pinned"
        older["scopes"]["SD"][0]["values"][CORE[0]] = "12"
        session.add(LocalRelease(id=older["id"], payload=older))
        session.flush()
        for release in (payload, older):
            query = {"release_id": release["id"]}
            assert client.get("/api/v1/workspace", params=query).json()["id"] == release["id"]
            assert client.get("/api/v1/regions", params=query).status_code == 200
            pinned = client.get("/api/v1/regions/11", params=query).json()
            assert pinned["release_id"] == release["id"]
            assert pinned["values"][CORE[0]] == release["scopes"]["SD"][0]["values"][CORE[0]]
        assert client.get("/api/v1/workspace?release_id=absent").status_code == 404
